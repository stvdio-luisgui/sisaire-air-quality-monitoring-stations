## SISAIRE Analysis Tool
# Luis Guillermo Pinilla Rodríguez

import pandas as pd
from tkinter import Tk, filedialog
from pathlib import Path
import re

from openpyxl.styles import (
    PatternFill,
    Font,
    Border,
    Side,
    Alignment
)

from openpyxl.utils import get_column_letter


# ============================================================
# DATABASE SELECTION AND LOADING
# ============================================================

ventana = Tk()
ventana.withdraw()

ruta = filedialog.askopenfilename(
    title="Seleccionar archivo CSV.GZ",
    filetypes=[
        ("Archivos CSV comprimidos", "*.csv.gz"),
        ("Todos los archivos", "*.*")
    ]
)

ventana.destroy()

if not ruta:
    raise ValueError("No se seleccionó ningún archivo.")

ruta = Path(ruta)

datos = pd.read_csv(
    ruta,
    compression="gzip",
    header=0,
    sep=";",
    encoding="cp1252"
)


# ============================================================
# CHECK REQUIRED COLUMNS
# ============================================================

columnas_requeridas = [
    "NOMBRE_FGDA",
    "NOMBRE_EST",
    "MSFL_CODE",
    "MED_CONCENTRACION_ESTANDAR"
]

columnas_faltantes = [
    columna
    for columna in columnas_requeridas
    if columna not in datos.columns
]

if columnas_faltantes:
    raise ValueError(
        "Faltan las siguientes columnas en la base de datos: "
        + ", ".join(columnas_faltantes)
    )


# ============================================================
# VALIDATE MED_CONCENTRACION_ESTANDAR
# ============================================================

# Preserve the original column before converting it to numeric.
# This allows comparison between the source values and the
# processed numeric values.

columna_original = datos[
    "MED_CONCENTRACION_ESTANDAR"
].copy()


# Count missing values before numeric conversion.

faltantes_antes = (
    columna_original
    .isna()
    .sum()
)

print(
    "\nVALORES FALTANTES ANTES "
    "DE LA CONVERSIÓN NUMÉRICA:"
)

print(faltantes_antes)


# Convert the measurement column to numeric.
#
# The original database may use commas as decimal separators.
# Example:
#
# "1,25" -> "1.25" -> 1.25
#
# Values that cannot be interpreted as numbers are converted
# to NaN using errors="coerce".

datos["MED_CONCENTRACION_ESTANDAR"] = pd.to_numeric(
    datos["MED_CONCENTRACION_ESTANDAR"]
    .astype("string")
    .str.replace(",", ".", regex=False),
    errors="coerce"
)


# Count missing values after conversion.

faltantes_despues = (
    datos["MED_CONCENTRACION_ESTANDAR"]
    .isna()
    .sum()
)

print(
    "\nVALORES FALTANTES DESPUÉS "
    "DE LA CONVERSIÓN NUMÉRICA:"
)

print(faltantes_despues)


# Identify values that existed in the original database
# but became NaN because they could not be converted
# to numeric values.

valores_no_convertibles = columna_original[
    columna_original.notna()
    & datos["MED_CONCENTRACION_ESTANDAR"].isna()
]

print(
    "\nVALORES QUE NO PUDIERON "
    "CONVERTIRSE A NÚMERO:"
)

print(
    len(valores_no_convertibles)
)

print(
    "\nVALORES NO CONVERTIBLES ENCONTRADOS:\n"
)


if len(valores_no_convertibles) > 0:

    conteo_no_convertibles = (
        valores_no_convertibles
        .astype("string")
        .map(repr)
        .value_counts()
    )

    print(
        conteo_no_convertibles.head(50)
    )

else:
    print("Ninguno")


# ============================================================
# GENERAL DATABASE INFORMATION
# ============================================================

print("\nArchivo cargado correctamente")

print("\nTAMAÑO DE LA BASE DE DATOS:")

print(
    "Filas:",
    datos.shape[0]
)

print(
    "Columnas:",
    datos.shape[1]
)


# Display the exact column names.
# repr() helps reveal hidden spaces or unusual characters.

print("\nNOMBRES DE LAS COLUMNAS:")

for columna in datos.columns:
    print(
        repr(columna)
    )


# Display pandas data types.

print("\nTIPOS DE DATOS:")

print(
    datos.dtypes
)


# Display the first ten complete records.

print("\nPRIMEROS 10 REGISTROS:")

print(
    datos.head(10).to_string(
        index=False,
        max_cols=None,
        max_colwidth=None
    )
)


# ============================================================
# ENVIRONMENTAL AUTHORITIES
# ============================================================

# Obtain all unique environmental authorities contained
# in NOMBRE_FGDA.

autoridades = (
    datos["NOMBRE_FGDA"]
    .dropna()
    .unique()
)

autoridades = sorted(
    autoridades
)

print(
    "\nAUTORIDADES AMBIENTALES ENCONTRADAS:\n"
)

for numero, autoridad in enumerate(
    autoridades,
    start=1
):
    print(
        f"{numero}. {autoridad}"
    )


# Count the number of records associated with each authority.

conteo_autoridades = (
    datos["NOMBRE_FGDA"]
    .value_counts(
        dropna=False
    )
)

print(
    "\nNÚMERO DE REGISTROS "
    "POR AUTORIDAD AMBIENTAL:\n"
)

print(
    conteo_autoridades
)


# ============================================================
# MONITORING STATIONS
# ============================================================

# Obtain all unique monitoring stations from NOMBRE_EST.

estaciones = (
    datos["NOMBRE_EST"]
    .dropna()
    .unique()
)

estaciones = sorted(
    estaciones
)

print(
    "\nESTACIONES ENCONTRADAS:\n"
)

for numero, estacion in enumerate(
    estaciones,
    start=1
):
    print(
        f"{numero}. {estacion}"
    )

print(
    "\nNÚMERO TOTAL DE ESTACIONES:",
    len(estaciones)
)


# ============================================================
# MEASURED VARIABLES
# ============================================================

# Obtain every unique measured variable stored in MSFL_CODE.

variables_medidas = (
    datos["MSFL_CODE"]
    .dropna()
    .unique()
)

variables_medidas = sorted(
    variables_medidas
)

print(
    "\nVARIABLES ENCONTRADAS:\n"
)

for numero, variable in enumerate(
    variables_medidas,
    start=1
):
    print(
        f"{numero}. {variable}"
    )


# Count the number of records associated with each variable.

conteo_variables = (
    datos["MSFL_CODE"]
    .value_counts(
        dropna=False
    )
)

print(
    "\nNÚMERO DE REGISTROS "
    "POR VARIABLE:\n"
)

print(
    conteo_variables
)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def nombre_archivo_seguro(nombre):

    # Convert the name to text and remove leading
    # and trailing spaces.

    nombre = str(nombre).strip()

    # Replace characters that Windows does not allow
    # in filenames.

    nombre = re.sub(
        r'[<>:"/\\|?*]',
        "_",
        nombre
    )

    return nombre


def nombre_hoja_seguro(
    nombre,
    hojas_usadas
):

    # Excel sheet names cannot contain:
    # [ ] : * ? / \
    #
    # These characters are replaced by underscores.

    nombre = re.sub(
        r'[\[\]:*?/\\]',
        "_",
        str(nombre)
    ).strip()

    # Assign a default name if the resulting name is empty.

    if not nombre:
        nombre = "HOJA"

    # Excel allows a maximum of 31 characters
    # in a worksheet name.

    nombre = nombre[:31]

    nombre_original = nombre

    contador = 2


    # Prevent duplicated worksheet names.

    while nombre in hojas_usadas:

        sufijo = f"_{contador}"

        nombre = (
            nombre_original[
                :31 - len(sufijo)
            ]
            + sufijo
        )

        contador += 1


    hojas_usadas.add(
        nombre
    )

    return nombre


# ============================================================
# PREFERRED VARIABLE ORDER
# ============================================================

# Main air-quality variables are placed first in the
# Excel workbook to keep related pollutants together.

orden_variables_principales = [
    "CO",
    "NO",
    "NO2",
    "NOX",
    "SO2",
    "O3",
    "PM10",
    "PM2.5"
]


# Keep only the preferred variables that actually exist
# in the complete database.

variables_ordenadas_global = [
    variable
    for variable in orden_variables_principales
    if variable in variables_medidas
]


# Add all other variables alphabetically afterwards.

otras_variables_global = sorted([
    variable
    for variable in variables_medidas
    if variable not in orden_variables_principales
])


variables_ordenadas_global = (
    variables_ordenadas_global
    + otras_variables_global
)


# ============================================================
# CREATE DATA PACKETS BY ENVIRONMENTAL AUTHORITY
# ============================================================

carpeta_autoridades = (
    ruta.parent
    / "carpeta_autoridades"
)

carpeta_autoridades.mkdir(
    exist_ok=True
)


archivos_autoridades_creados = []

numero_archivos = 0


# Split the complete database according to NOMBRE_FGDA.
#
# Each resulting file contains all records and all stations
# belonging to one environmental authority.

for autoridad, grupo in datos.groupby(
    "NOMBRE_FGDA",
    dropna=False,
    sort=False
):


    if pd.isna(autoridad):

        nombre_autoridad_archivo = (
            "SIN_NOMBRE_FGDA"
        )

    else:

        nombre_autoridad_archivo = (
            nombre_archivo_seguro(
                autoridad
            )
        )


    archivo_salida = (
        carpeta_autoridades
        / f"{nombre_autoridad_archivo}.csv.gz"
    )


    grupo.to_csv(
        archivo_salida,
        sep=";",
        index=False,
        encoding="cp1252",
        compression="gzip"
    )


    archivos_autoridades_creados.append(
        archivo_salida
    )


    numero_archivos += 1


    # Count the number of unique monitoring stations
    # belonging to the current authority.

    numero_estaciones_grupo = (
        grupo["NOMBRE_EST"]
        .dropna()
        .nunique()
    )


    print(
        f"{numero_archivos}. "
        f"{nombre_autoridad_archivo}: "
        f"{len(grupo):,} registros - "
        f"{numero_estaciones_grupo} estaciones"
    )


print(
    "\nNúmero de archivos "
    "de autoridades creados:",
    numero_archivos
)

print(
    "\nArchivos guardados en:"
)

print(
    carpeta_autoridades
)


# ============================================================
# OUTPUT FOLDERS
# ============================================================

# Folder containing the Excel workbook generated
# for each environmental authority.

carpeta_excel = (
    ruta.parent
    / "excel_autoridades"
)

carpeta_excel.mkdir(
    exist_ok=True
)


# Folder containing the summary databases.

carpeta_resumenes = (
    ruta.parent
    / "bases_resumen"
)

carpeta_resumenes.mkdir(
    exist_ok=True
)


# ============================================================
# LISTS USED TO BUILD SUMMARY DATABASES
# ============================================================

# One record per environmental authority.

resumen_autoridades_lista = []


# One record per environmental authority and MSFL variable.

resumen_msfl_lista = []


# One record per authority-station combination,
# including detailed variable statistics.

relacion_autoridad_estacion_lista = []


# ============================================================
# PROCESS ALL ENVIRONMENTAL AUTHORITIES
# ============================================================

archivos_autoridades = sorted(
    archivos_autoridades_creados
)

print(
    "\nPROCESANDO TODAS "
    "LAS AUTORIDADES AMBIENTALES"
)

print(
    "Autoridades encontradas:",
    len(archivos_autoridades)
)


# Process one authority at a time.
#
# This avoids trying to generate all authority workbooks
# simultaneously.

for numero_autoridad, archivo_autoridad in enumerate(
    archivos_autoridades,
    start=1
):


    print(
        f"\nProcesando autoridad "
        f"{numero_autoridad}/"
        f"{len(archivos_autoridades)}:"
    )

    print(
        archivo_autoridad.name
    )


    # ========================================================
    # LOAD ENVIRONMENTAL AUTHORITY DATA PACKET
    # ========================================================

    datos_autoridad = pd.read_csv(
        archivo_autoridad,
        compression="gzip",
        header=0,
        sep=";",
        encoding="cp1252"
    )


    # The authority packet has already been processed,
    # but explicitly ensure that the measurement column
    # is numeric.

    datos_autoridad[
        "MED_CONCENTRACION_ESTANDAR"
    ] = pd.to_numeric(
        datos_autoridad[
            "MED_CONCENTRACION_ESTANDAR"
        ],
        errors="coerce"
    )


    # ========================================================
    # IDENTIFY ENVIRONMENTAL AUTHORITY
    # ========================================================

    autoridades_archivo = (
        datos_autoridad[
            "NOMBRE_FGDA"
        ]
        .dropna()
        .unique()
    )


    if len(
        autoridades_archivo
    ) > 0:

        nombre_autoridad = (
            autoridades_archivo[0]
        )

    else:

        nombre_autoridad = (
            "SIN_NOMBRE_FGDA"
        )


    # ========================================================
    # VARIABLES AVAILABLE FOR THIS AUTHORITY
    # ========================================================

    variables_autoridad = sorted(
        datos_autoridad[
            "MSFL_CODE"
        ]
        .dropna()
        .unique()
    )


    # Preserve the preferred global order, but include only
    # variables actually available for this authority.

    variables_autoridad_ordenadas = [
        variable
        for variable in variables_ordenadas_global
        if variable in variables_autoridad
    ]


    # ========================================================
    # GENERAL AUTHORITY SUMMARY
    # ========================================================

    # Total number of database records belonging
    # to this environmental authority.

    total_registros = (
        len(
            datos_autoridad
        )
    )


    # Number of unique monitoring stations.

    numero_estaciones = (
        datos_autoridad[
            "NOMBRE_EST"
        ]
        .dropna()
        .nunique()
    )


    # Number of variables available in the complete database.

    total_variables_base = (
        len(
            variables_medidas
        )
    )


    # Number of variables measured by this authority.

    total_variables_medidas = (
        len(
            variables_autoridad
        )
    )


    # Count actual numeric zero measurements.

    total_ceros = (
        datos_autoridad[
            "MED_CONCENTRACION_ESTANDAR"
        ] == 0
    ).sum()


    # Count missing measurements.

    total_faltantes = (
        datos_autoridad[
            "MED_CONCENTRACION_ESTANDAR"
        ]
        .isna()
        .sum()
    )


    # Store one general summary record for this authority.

    resumen_autoridad = {

        "NOMBRE_FGDA":
            nombre_autoridad,

        "ARCHIVO_AUTORIDAD":
            archivo_autoridad.name,

        "TOTAL_REGISTROS":
            total_registros,

        "NUMERO_ESTACIONES":
            numero_estaciones,

        "VARIABLES_BASE_DATOS":
            total_variables_base,

        "VARIABLES_MEDIDAS":
            total_variables_medidas,

        "VARIABLES_NO_MEDIDAS":
            (
                total_variables_base
                - total_variables_medidas
            ),

        "TOTAL_CEROS":
            total_ceros,

        "TOTAL_FALTANTES":
            total_faltantes
    }


    resumen_autoridades_lista.append(
        resumen_autoridad
    )


    # ========================================================
    # COMPLETE SUMMARY BY MONITORING STATION
    # ========================================================

    # Create a temporary copy for summary calculations.

    datos_resumen = (
        datos_autoridad.copy()
    )


    # Records without a station name are retained and
    # explicitly labelled instead of being discarded.

    datos_resumen[
        "NOMBRE_EST"
    ] = (
        datos_resumen[
            "NOMBRE_EST"
        ]
        .fillna(
            "SIN_NOMBRE_EST"
        )
    )


    # Calculate total records, zero values and missing values
    # for each monitoring station.

    resumen_totales_estacion = (
        datos_resumen
        .groupby(
            "NOMBRE_EST"
        )
        .agg(

            TOTAL_REGISTROS=(
                "MED_CONCENTRACION_ESTANDAR",
                "size"
            ),

            TOTAL_CEROS=(
                "MED_CONCENTRACION_ESTANDAR",
                lambda x: (
                    x == 0
                ).sum()
            ),

            TOTAL_FALTANTES=(
                "MED_CONCENTRACION_ESTANDAR",
                lambda x:
                    x.isna().sum()
            )
        )
    )


    # ========================================================
    # SUMMARY BY STATION AND VARIABLE
    # ========================================================

    # Calculate records, zeros and missing values
    # for every station-variable combination.

    resumen_estacion_variable = (
        datos_resumen
        .groupby(
            [
                "NOMBRE_EST",
                "MSFL_CODE"
            ]
        )
        .agg(

            REGISTROS=(
                "MED_CONCENTRACION_ESTANDAR",
                "size"
            ),

            CEROS=(
                "MED_CONCENTRACION_ESTANDAR",
                lambda x: (
                    x == 0
                ).sum()
            ),

            FALTANTES=(
                "MED_CONCENTRACION_ESTANDAR",
                lambda x:
                    x.isna().sum()
            )
        )
    )


    # ========================================================
    # RECORDS BY VARIABLE
    # ========================================================

    # Convert the variable dimension from rows into columns.
    #
    # Example:
    #
    # CO_REGISTROS
    # NO2_REGISTROS
    # PM10_REGISTROS
    # ...

    registros_por_variable = (
        resumen_estacion_variable[
            "REGISTROS"
        ]
        .unstack(
            fill_value=0
        )
        .reindex(
            columns=variables_ordenadas_global,
            fill_value=0
        )
    )


    registros_por_variable.columns = [
        f"{variable}_REGISTROS"
        for variable
        in registros_por_variable.columns
    ]


    # ========================================================
    # ZERO VALUES BY VARIABLE
    # ========================================================

    ceros_por_variable = (
        resumen_estacion_variable[
            "CEROS"
        ]
        .unstack(
            fill_value=0
        )
        .reindex(
            columns=variables_ordenadas_global,
            fill_value=0
        )
    )


    ceros_por_variable.columns = [
        f"{variable}_CEROS"
        for variable
        in ceros_por_variable.columns
    ]


    # ========================================================
    # MISSING VALUES BY VARIABLE
    # ========================================================

    faltantes_por_variable = (
        resumen_estacion_variable[
            "FALTANTES"
        ]
        .unstack(
            fill_value=0
        )
        .reindex(
            columns=variables_ordenadas_global,
            fill_value=0
        )
    )


    faltantes_por_variable.columns = [
        f"{variable}_FALTANTES"
        for variable
        in faltantes_por_variable.columns
    ]


    # ========================================================
    # MERGE STATION SUMMARY TABLES
    # ========================================================

    resumen_estaciones_autoridad = (
        resumen_totales_estacion
        .join(
            registros_por_variable
        )
        .join(
            ceros_por_variable
        )
        .join(
            faltantes_por_variable
        )
        .fillna(0)
        .reset_index()
    )


    # ========================================================
    # ORDER SUMMARY COLUMNS
    # ========================================================

    # General station totals are placed first.

    columnas_ordenadas = [
        "NOMBRE_EST",
        "TOTAL_REGISTROS",
        "TOTAL_CEROS",
        "TOTAL_FALTANTES"
    ]


    # Each variable is then represented by a group
    # of three consecutive columns:
    #
    # REGISTROS
    # CEROS
    # FALTANTES

    for variable in variables_ordenadas_global:

        columnas_ordenadas.extend([
            f"{variable}_REGISTROS",
            f"{variable}_CEROS",
            f"{variable}_FALTANTES"
        ])


    resumen_estaciones_autoridad = (
        resumen_estaciones_autoridad[
            columnas_ordenadas
        ]
        .sort_values(
            "NOMBRE_EST"
        )
        .reset_index(
            drop=True
        )
    )


    # Identify all numeric summary columns.

    columnas_numericas = [
        columna
        for columna
        in resumen_estaciones_autoridad.columns
        if columna != "NOMBRE_EST"
    ]


    # Store all count columns as integer values.

    resumen_estaciones_autoridad[
        columnas_numericas
    ] = (
        resumen_estaciones_autoridad[
            columnas_numericas
        ]
        .fillna(0)
        .astype(int)
    )


    # ========================================================
    # STORE AUTHORITY-STATION SUMMARY IN PYTHON
    # ========================================================

    # Each station becomes one row in the general
    # authority-station summary database.

    for _, fila_estacion in (
        resumen_estaciones_autoridad.iterrows()
    ):

        registro_estacion = {

            "NOMBRE_FGDA":
                nombre_autoridad
        }


        registro_estacion.update(
            fila_estacion.to_dict()
        )


        relacion_autoridad_estacion_lista.append(
            registro_estacion
        )


    # ========================================================
    # AUTHORITY TOTAL ROW
    # ========================================================

    # Add a final row containing the sum of all stations.

    fila_total = {

        "NOMBRE_EST":
            "TOTAL AUTORIDAD"
    }


    for columna in columnas_numericas:

        fila_total[columna] = (
            resumen_estaciones_autoridad[
                columna
            ]
            .sum()
        )


    resumen_estaciones_excel = pd.concat(
        [
            resumen_estaciones_autoridad,
            pd.DataFrame(
                [fila_total]
            )
        ],
        ignore_index=True
    )


    # ========================================================
    # VARIABLE SUMMARY BY AUTHORITY
    # ========================================================

    # This summary remains available in Python even though
    # the Excel summary sheet uses the wider station matrix.

    conteo_autoridad_por_variable = (
        datos_autoridad
        .groupby(
            "MSFL_CODE"
        )
        .agg(

            REGISTROS=(
                "MED_CONCENTRACION_ESTANDAR",
                "size"
            ),

            CEROS=(
                "MED_CONCENTRACION_ESTANDAR",
                lambda x: (
                    x == 0
                ).sum()
            ),

            FALTANTES=(
                "MED_CONCENTRACION_ESTANDAR",
                lambda x:
                    x.isna().sum()
            )
        )
    )


    # Build a table using all variables from the complete
    # SISAIRE database so variables not measured by an
    # authority are still explicitly represented.

    resumen_variables_autoridad = (
        pd.DataFrame(
            index=variables_ordenadas_global
        )
    )


    resumen_variables_autoridad[
        "MEDIDA"
    ] = (
        resumen_variables_autoridad
        .index
        .isin(
            variables_autoridad
        )
    )


    resumen_variables_autoridad[
        "MEDIDA"
    ] = (
        resumen_variables_autoridad[
            "MEDIDA"
        ]
        .map({
            True: "SÍ",
            False: "NO"
        })
    )


    resumen_variables_autoridad = (
        resumen_variables_autoridad
        .join(
            conteo_autoridad_por_variable
        )
    )


    resumen_variables_autoridad[
        [
            "REGISTROS",
            "CEROS",
            "FALTANTES"
        ]
    ] = (
        resumen_variables_autoridad[
            [
                "REGISTROS",
                "CEROS",
                "FALTANTES"
            ]
        ]
        .fillna(0)
        .astype(int)
    )


    resumen_variables_autoridad.index.name = (
        "MSFL_CODE"
    )


    resumen_variables_autoridad = (
        resumen_variables_autoridad
        .reset_index()
    )


    resumen_variables_autoridad.insert(
        0,
        "NOMBRE_FGDA",
        nombre_autoridad
    )


    resumen_msfl_lista.append(
        resumen_variables_autoridad.copy()
    )


    # ========================================================
    # GENERAL INFORMATION FOR THE EXCEL SUMMARY SHEET
    # ========================================================

    resumen_autoridad_excel = (
        pd.DataFrame({

            "INFORMACIÓN DE LA AUTORIDAD": [

                "Autoridad ambiental",

                "Archivo de la autoridad",

                "Total de registros",

                "Número de estaciones",

                "Variables en la base de datos completa",

                "Variables medidas por la autoridad",

                "Variables no medidas por la autoridad",

                "Total de valores cero",

                "Total de valores faltantes"
            ],

            "VALOR": [

                nombre_autoridad,

                archivo_autoridad.name,

                total_registros,

                numero_estaciones,

                total_variables_base,

                total_variables_medidas,

                (
                    total_variables_base
                    - total_variables_medidas
                ),

                total_ceros,

                total_faltantes
            ]
        })
    )


    # ========================================================
    # CREATE EXCEL WORKBOOK
    # ========================================================

    nombre_excel = (
        archivo_autoridad.name
        .removesuffix(
            ".csv.gz"
        )
        + ".xlsx"
    )


    archivo_excel = (
        carpeta_excel
        / nombre_excel
    )


    with pd.ExcelWriter(
        archivo_excel,
        engine="openpyxl"
    ) as escritor:


        # ====================================================
        # WRITE GENERAL AUTHORITY INFORMATION
        # ====================================================

        resumen_autoridad_excel.to_excel(
            escritor,
            sheet_name="RESUMEN",
            index=False,
            startrow=0
        )


        # ====================================================
        # WRITE STATION-VARIABLE SUMMARY MATRIX
        # ====================================================

        fila_estaciones = (
            len(
                resumen_autoridad_excel
            )
            + 3
        )


        resumen_estaciones_excel.to_excel(
            escritor,
            sheet_name="RESUMEN",
            index=False,
            startrow=fila_estaciones
        )


        # ====================================================
        # FORMAT EXCEL SUMMARY SHEET
        # ====================================================

        hoja_resumen = (
            escritor.book[
                "RESUMEN"
            ]
        )


        # Excel row containing the summary table headers.

        fila_encabezado = (
            fila_estaciones
            + 1
        )


        # Excel row containing the final authority total.

        fila_total_excel = (
            fila_encabezado
            + len(
                resumen_estaciones_excel
            )
        )


        ultima_columna = (
            hoja_resumen.max_column
        )


        # ----------------------------------------------------
        # COLOURS
        # ----------------------------------------------------

        # Light grey for station names.

        relleno_nombre = PatternFill(
            fill_type="solid",
            fgColor="E7E6E6"
        )


        # Blue shades for the general TOTAL block.

        relleno_total_encabezado = PatternFill(
            fill_type="solid",
            fgColor="B4C6E7"
        )


        relleno_total_datos = PatternFill(
            fill_type="solid",
            fgColor="EAF0F8"
        )


        # Two alternating green shades are used for
        # successive variable groups.

        relleno_grupo_1_encabezado = PatternFill(
            fill_type="solid",
            fgColor="C6E0B4"
        )


        relleno_grupo_1_datos = PatternFill(
            fill_type="solid",
            fgColor="F0F7EC"
        )


        relleno_grupo_2_encabezado = PatternFill(
            fill_type="solid",
            fgColor="D9EAD3"
        )


        relleno_grupo_2_datos = PatternFill(
            fill_type="solid",
            fgColor="F7FAF5"
        )


        # Yellow highlight for the final authority total row.

        relleno_total_autoridad = PatternFill(
            fill_type="solid",
            fgColor="FFF2CC"
        )


        # ----------------------------------------------------
        # BORDERS
        # ----------------------------------------------------

        # Medium grey border separates each three-column
        # variable block.

        borde_grupo = Side(
            style="medium",
            color="808080"
        )


        # Dark border separates the final total row.

        borde_total = Side(
            style="medium",
            color="000000"
        )


        # ----------------------------------------------------
        # STATION NAME COLUMN
        # ----------------------------------------------------

        # Apply a light grey background to the station-name
        # column to visually separate row labels from values.

        for fila in range(
            fila_encabezado,
            fila_total_excel + 1
        ):

            celda = hoja_resumen.cell(
                row=fila,
                column=1
            )

            celda.fill = (
                relleno_nombre
            )


        # Make the NOMBRE_EST header bold.

        hoja_resumen.cell(
            row=fila_encabezado,
            column=1
        ).font = Font(
            bold=True
        )


        # ----------------------------------------------------
        # GENERAL TOTAL BLOCK
        # ----------------------------------------------------

        # Columns 2 to 4 correspond to:
        #
        # TOTAL_REGISTROS
        # TOTAL_CEROS
        # TOTAL_FALTANTES

        for columna in range(
            2,
            5
        ):

            celda_encabezado = (
                hoja_resumen.cell(
                    row=fila_encabezado,
                    column=columna
                )
            )

            celda_encabezado.fill = (
                relleno_total_encabezado
            )

            celda_encabezado.font = Font(
                bold=True
            )

            celda_encabezado.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )


            # Apply a lighter version of the same colour
            # to the data cells.

            for fila in range(
                fila_encabezado + 1,
                fila_total_excel
            ):

                hoja_resumen.cell(
                    row=fila,
                    column=columna
                ).fill = (
                    relleno_total_datos
                )


        # Add a vertical border after the TOTAL block.

        for fila in range(
            fila_encabezado,
            fila_total_excel + 1
        ):

            hoja_resumen.cell(
                row=fila,
                column=4
            ).border = Border(
                right=borde_grupo
            )


        # ----------------------------------------------------
        # VARIABLE BLOCKS
        # ----------------------------------------------------

        # Every measured variable occupies three columns:
        #
        # VARIABLE_REGISTROS
        # VARIABLE_CEROS
        # VARIABLE_FALTANTES
        #
        # Alternating colours make adjacent variable groups
        # easier to distinguish.

        columna_inicio = 5

        numero_grupo = 0


        while (
            columna_inicio
            <= ultima_columna
        ):


            columna_fin = min(
                columna_inicio + 2,
                ultima_columna
            )


            if numero_grupo % 2 == 0:

                relleno_encabezado = (
                    relleno_grupo_1_encabezado
                )

                relleno_datos = (
                    relleno_grupo_1_datos
                )

            else:

                relleno_encabezado = (
                    relleno_grupo_2_encabezado
                )

                relleno_datos = (
                    relleno_grupo_2_datos
                )


            for columna in range(
                columna_inicio,
                columna_fin + 1
            ):


                celda_encabezado = (
                    hoja_resumen.cell(
                        row=fila_encabezado,
                        column=columna
                    )
                )


                celda_encabezado.fill = (
                    relleno_encabezado
                )


                celda_encabezado.font = Font(
                    bold=True
                )


                celda_encabezado.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )


                # Apply the corresponding lighter colour
                # to the data section.

                for fila in range(
                    fila_encabezado + 1,
                    fila_total_excel
                ):

                    hoja_resumen.cell(
                        row=fila,
                        column=columna
                    ).fill = (
                        relleno_datos
                    )


            # Add a vertical border at the end of each
            # variable block.

            for fila in range(
                fila_encabezado,
                fila_total_excel + 1
            ):

                hoja_resumen.cell(
                    row=fila,
                    column=columna_fin
                ).border = Border(
                    right=borde_grupo
                )


            columna_inicio += 3

            numero_grupo += 1


        # ----------------------------------------------------
        # AUTHORITY TOTAL ROW
        # ----------------------------------------------------

        # Highlight the final TOTAL AUTORIDAD row using
        # bold text, a yellow background and a top border.

        for columna in range(
            1,
            ultima_columna + 1
        ):


            celda = hoja_resumen.cell(
                row=fila_total_excel,
                column=columna
            )


            celda.font = Font(
                bold=True
            )


            celda.fill = (
                relleno_total_autoridad
            )


            celda.border = Border(
                top=borde_total
            )


        # ----------------------------------------------------
        # ALIGNMENT
        # ----------------------------------------------------

        # Station names are left-aligned.

        for fila in range(
            fila_encabezado,
            fila_total_excel + 1
        ):

            hoja_resumen.cell(
                row=fila,
                column=1
            ).alignment = Alignment(
                horizontal="left",
                vertical="center"
            )


            # Numeric summary values are centred.

            for columna in range(
                2,
                ultima_columna + 1
            ):

                hoja_resumen.cell(
                    row=fila,
                    column=columna
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )


        # ----------------------------------------------------
        # COLUMN WIDTHS
        # ----------------------------------------------------

        # Station names need a wider column.

        hoja_resumen.column_dimensions[
            "A"
        ].width = 45


        # Summary numeric columns use a consistent width.

        for columna in range(
            2,
            ultima_columna + 1
        ):

            letra_columna = (
                get_column_letter(
                    columna
                )
            )

            hoja_resumen.column_dimensions[
                letra_columna
            ].width = 16


        # ----------------------------------------------------
        # HEADER ROW HEIGHT
        # ----------------------------------------------------

        # Increase header height because some variable names
        # may wrap onto multiple lines.

        hoja_resumen.row_dimensions[
            fila_encabezado
        ].height = 35


        # ====================================================
        # VARIABLE DATA SHEETS
        # ====================================================

        hojas_usadas = {
            "RESUMEN"
        }


        # Create one worksheet for each variable available
        # for this environmental authority.

        for variable in variables_autoridad_ordenadas:


            datos_variable = (
                datos_autoridad[
                    datos_autoridad[
                        "MSFL_CODE"
                    ] == variable
                ]
            )


            # Excel supports a maximum of 1,048,576 rows
            # per worksheet. One row is reserved for the header.

            maximo_filas_datos = (
                1_048_575
            )


            # Calculate how many worksheet parts are required.

            numero_partes = max(
                1,
                (
                    len(datos_variable)
                    + maximo_filas_datos
                    - 1
                )
                // maximo_filas_datos
            )


            for parte in range(
                numero_partes
            ):


                inicio = (
                    parte
                    * maximo_filas_datos
                )


                fin = (
                    inicio
                    + maximo_filas_datos
                )


                datos_parte = (
                    datos_variable
                    .iloc[
                        inicio:fin
                    ]
                )


                # Use the variable name directly when the
                # complete dataset fits into one worksheet.

                if numero_partes == 1:

                    nombre_base_hoja = (
                        str(variable)
                    )

                # Otherwise create multiple numbered sheets,
                # for example PM10_1, PM10_2, etc.

                else:

                    nombre_base_hoja = (
                        f"{variable}_{parte + 1}"
                    )


                nombre_hoja = (
                    nombre_hoja_seguro(
                        nombre_base_hoja,
                        hojas_usadas
                    )
                )


                datos_parte.to_excel(
                    escritor,
                    sheet_name=nombre_hoja,
                    index=False
                )


    print(
        "Excel creado:",
        archivo_excel.name
    )


# ============================================================
# CREATE SUMMARY DATABASES IN PYTHON
# ============================================================

# Database with one row per environmental authority.

resumen_autoridades_db = (
    pd.DataFrame(
        resumen_autoridades_lista
    )
)


# Database with one row for every
# authority-variable combination.

if resumen_msfl_lista:

    resumen_msfl_db = (
        pd.concat(
            resumen_msfl_lista,
            ignore_index=True
        )
    )

else:

    resumen_msfl_db = (
        pd.DataFrame()
    )


# Database with one row per authority-station combination
# and detailed statistics for every measured variable.

relacion_autoridad_estacion_db = (
    pd.DataFrame(
        relacion_autoridad_estacion_lista
    )
)


# ============================================================
# SAVE SUMMARY DATABASES
# ============================================================

archivo_resumen_autoridades = (
    carpeta_resumenes
    / "resumen_autoridades.csv.gz"
)


archivo_resumen_msfl = (
    carpeta_resumenes
    / "resumen_msfl.csv.gz"
)


archivo_relacion_autoridad_estacion = (
    carpeta_resumenes
    / "relacion_autoridad_estacion.csv.gz"
)


# Save general authority summary.

resumen_autoridades_db.to_csv(
    archivo_resumen_autoridades,
    sep=";",
    index=False,
    encoding="cp1252",
    compression="gzip"
)


# Save authority-variable summary.

resumen_msfl_db.to_csv(
    archivo_resumen_msfl,
    sep=";",
    index=False,
    encoding="cp1252",
    compression="gzip"
)


# Save authority-station-variable summary.

relacion_autoridad_estacion_db.to_csv(
    archivo_relacion_autoridad_estacion,
    sep=";",
    index=False,
    encoding="cp1252",
    compression="gzip"
)


# ============================================================
# FINAL RESULTS
# ============================================================

print(
    "\nPROCESO COMPLETADO"
)


print(
    "\nAutoridades procesadas:",
    len(
        resumen_autoridades_db
    )
)


print(
    "\nRESUMEN DE AUTORIDADES:"
)

print(
    resumen_autoridades_db.head()
)


print(
    "\nRESUMEN AUTORIDAD - ESTACIÓN - VARIABLES:"
)

print(
    relacion_autoridad_estacion_db.head()
)


print(
    "\nRESUMEN DE VARIABLES MSFL:"
)

print(
    resumen_msfl_db.head()
)


print(
    "\nArchivos por autoridad guardados en:"
)

print(
    carpeta_autoridades
)


print(
    "\nArchivos Excel guardados en:"
)

print(
    carpeta_excel
)


print(
    "\nBases de resumen guardadas en:"
)

print(
    carpeta_resumenes
)